"""Thẻ vào PTN — nhập / xuất Excel qua API.

1. Nhập luôn XEM TRƯỚC được (dry_run) mà không ghi gì.
2. Đối chiếu theo MSSV: có thì cập nhật, chưa có thì tạo — nhập lại lần hai không
   sinh bản trùng.
3. Chỉ người quản trị thẻ (office/admin) được nhập; ai đọc được thì xuất được.
"""
import io

from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

from app.models.lab_access import LabAccessCard
from app.tests.conftest import requires_db

_BASE = "/api/v1/lab-access-cards"


def _xlsx(rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "QUÝ I.2026"
    ws["A1"] = "DANH SÁCH ĐÃ ĐƯỢC CẤP THẺ VÀO PTN QUÝ I.2026 (TỔNG)"
    ws.append([])
    ws.append(["STT", "HỌ", "TÊN", "LỚP", "MSSV", "EMAIL", "ĐĂNG KÝ SỬ DỤNG PTN",
               "MỤC ĐÍCH SỬ DỤNG", "GIÁO VIÊN HƯỚNG DẪN", "THỜI GIAN", "GHI CHÚ"])
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _post(client, content: bytes, dry_run: bool):
    return client.post(
        f"{_BASE}/import",
        params={"dry_run": str(dry_run).lower()},
        files={"file": ("the.xlsx", content, "application/octet-stream")},
    )


def _count(db, code: str) -> int:
    return db.execute(
        select(func.count()).select_from(LabAccessCard).where(LabAccessCard.student_code == code)
    ).scalar_one()


@requires_db
def test_import_preview_then_commit_then_update(client, db, as_role):
    as_role("office")
    content = _xlsx([
        [1, "Test", "Một", "DH23SHA", 99900001, "a@x", "RIBE 306", "KLTN", "TS. A", "06/4/26-06/7/26", None],
        [2, "Test", "Hai", "DH23SHA", 99900002, None, None, "Rèn nghề", "TS. A", "06/4/26-06/7/26", "GH"],
    ])

    res = _post(client, content, dry_run=True)
    assert res.status_code == 200, res.text
    data = res.json()["data"]
    assert (data["dry_run"], data["created"], data["errors"]) == (True, 2, [])
    assert _count(db, "99900001") == 0

    res = _post(client, content, dry_run=False)
    assert res.json()["data"]["created"] == 2
    assert _count(db, "99900001") == 1

    # Lần hai: một dòng đổi phòng → update, dòng kia giữ nguyên; không tạo trùng.
    changed = _xlsx([
        [1, "Test", "Một", "DH23SHA", 99900001, "a@x", "RIBE 999", "KLTN", "TS. A", "06/4/26-06/7/26", None],
        [2, "Test", "Hai", "DH23SHA", 99900002, None, None, "Rèn nghề", "TS. A", "06/4/26-06/7/26", "GH"],
    ])
    data = _post(client, changed, dry_run=False).json()["data"]
    assert (data["created"], data["updated"], data["unchanged"]) == (0, 1, 1)
    assert _count(db, "99900001") == 1
    card = db.execute(
        select(LabAccessCard).where(LabAccessCard.student_code == "99900001")
    ).scalar_one()
    assert card.room == "RIBE 999"


@requires_db
def test_import_forbidden_for_readers(client, as_role):
    as_role("qms")
    res = _post(client, _xlsx([]), dry_run=True)
    assert res.status_code == 403


@requires_db
def test_export_xlsx(client, as_role):
    as_role("office")
    _post(client, _xlsx([
        [1, "Xuất", "Thử", "DH23SHA", 99900003, None, "BIO 203", "KLTN", "ThS. B", "24.3.2026-24.9.2026", None],
    ]), dry_run=False)
    res = client.get(f"{_BASE}/export.xlsx", params={"q": "99900003", "quarter": 2, "year": 2026})
    assert res.status_code == 200, res.text
    assert 'filename="the-vao-ptn-quy-II-2026.xlsx"' in res.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(res.content))
    assert wb.sheetnames == ["QUÝ II.2026", "B"]
    ws = wb.worksheets[0]
    assert ws["A1"].value == "DANH SÁCH ĐÃ ĐƯỢC CẤP THẺ VÀO PTN QUÝ II.2026 (TỔNG)"
    assert [c.value for c in ws[4]][:5] == [1, "Xuất", "Thử", "DH23SHA", 99900003]
