"""Xuất Excel danh sách Nghiên cứu & Đào tạo, lọc theo khoảng thời gian.

BA NHÓM BẤT BIẾN
1. Lọc thời gian ĐÚNG với từng loại ngày nghiệp vụ — mỗi mục một trường khác nhau
   (ngày ký hợp đồng, ngày thực hiện, ngày cấp…), và ba mục chỉ có NĂM nên khoảng
   ngày phải quy về khoảng năm.
2. Bản ghi thiếu dữ liệu thời gian bị loại nhưng ĐƯỢC ĐẾM — âm thầm bỏ bớt dòng là
   cách nhanh nhất làm hỏng một báo cáo nộp lên trên.
3. File xuất ra chỉ chứa đúng phần người dùng được xem: phạm vi do chính hàm list_*
   quyết định, nên KTV chỉ xuất được bản ghi của mình.
"""
import io
import re
from datetime import date, timedelta

import pytest

from app.tests.conftest import requires_db

_BASE = "/api/v1/research-exports"
_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

ALL_KINDS = (
    "research-projects", "publications", "research-contracts", "community-services",
    "student-mentorships", "teaching-courses", "training-certificates", "staff-activities",
    # m47 — tách "Bài báo & Sáng chế" làm hai mục có bộ cột riêng. Mục gộp
    # 'publications' ở trên GIỮ NGUYÊN cho báo cáo tổng kết năm.
    "papers", "patents",
)


@pytest.fixture(autouse=True)
def _reset_rate_limit():
    """Xoá bộ đếm rate-limit trước mỗi test.

    `rate_limit` đếm theo ĐỊA CHỈ IP, mà cả bộ test chạy qua một TestClient duy nhất
    nên mọi test dùng chung một bộ đếm. Chạy riêng file này thì đủ hạn mức, chạy cả
    bộ thì tràn — tức là test đỏ vì thứ tự chạy chứ không vì lỗi thật. Ở đây ta kiểm
    logic xuất file, không kiểm bộ giới hạn tần suất.
    """
    from app.core.redis_client import get_redis

    r = get_redis()
    for k in r.scan_iter("ratelimit:research-export:*"):
        r.delete(k)
    yield


def _sheet(res):
    """Đọc lại file vừa tải để khẳng định trên NỘI DUNG THẬT, không chỉ mã 200."""
    from openpyxl import load_workbook

    assert res.headers["content-type"].startswith(_XLSX), res.headers["content-type"]
    wb = load_workbook(io.BytesIO(res.content))
    ws = wb.active
    return [[c.value for c in row] for row in ws.iter_rows()]


def _meta(rows: list, label: str):
    """Đọc một con số / mẩu thông tin từ KHỐI ĐẦU TRANG.

    m48 gộp khối này lại: thay vì 6 dòng nhãn/giá trị rời ở cột A và B, nay là bốn
    dòng câu văn đã gộp ô. Test vẫn hỏi theo NHÃN NGHIỆP VỤ như cũ ("Số bản ghi"),
    còn việc nhãn đó nằm ở đâu trong file là chuyện của helper này — nhờ vậy đổi bố
    cục lần sau chỉ phải sửa một chỗ.
    """
    head = " ".join(str(r[0]) for r in rows[:5] if r and r[0])

    if label == "Số bản ghi":
        m = re.search(r"(\d+) bản ghi", head)
        return int(m.group(1)) if m else None
    if label == "Khoảng thời gian":
        m = re.search(r"Kỳ báo cáo: (.+?)\s+·", head)
        return m.group(1).strip() if m else None
    if label == "Tiêu chí lọc":
        # Tiêu chí lọc gồm TÊN TRƯỜNG và, với mục chỉ lưu năm, cả lời cảnh báo rằng
        # khoảng ngày đã bị quy về khoảng năm — hai mẩu này thuộc cùng một ý.
        m = re.search(r"Lọc theo: (.+)$", head)
        if not m:
            return None
        phan = [p.strip() for p in re.split(r"\s+·\s+", m.group(1)) if p.strip()]
        return " · ".join(p for p in phan if not re.match(r"^\d+ bản ghi", p))
    if label.startswith("Bị loại"):
        m = re.search(r"(\d+) bản ghi bị loại", head)
        return int(m.group(1)) if m else None
    return None


def _data_rows(rows: list) -> list:
    """Chỉ các dòng DỮ LIỆU, đã bỏ cột STT.

    Bỏ STT để test vẫn đánh chỉ số cột như trước m48 (r[0] = cột dữ liệu đầu tiên);
    bỏ dòng TỔNG ở cuối vì nó không phải một bản ghi.
    """
    from app.services.research_export_service import _HEADER_ROW

    data = rows[_HEADER_ROW:]  # bỏ khối đầu trang + dòng tiêu đề cột
    if data and data[-1] and data[-1][0] == "TỔNG":
        data = data[:-1]
    return [r[1:] for r in data]


@requires_db
class TestDanhMucMucXuat:
    def test_liet_ke_du_moi_muc(self, client, as_role):
        as_role("leader")
        res = client.get(_BASE)
        assert res.status_code == 200, res.text
        kinds = {k["kind"] for k in res.json()["data"]}
        assert kinds == set(ALL_KINDS)

    def test_moi_muc_noi_ro_loc_theo_truong_nao(self, client, as_role):
        """Người dùng phải biết khoảng thời gian đang lọc theo cái gì."""
        as_role("leader")
        by_kind = {k["kind"]: k for k in client.get(_BASE).json()["data"]}
        assert by_kind["research-contracts"]["filter_field"] == "Ngày ký"
        assert by_kind["community-services"]["filter_field"] == "Ngày thực hiện"
        # Ba mục chỉ có năm phải được đánh dấu để giao diện chú thích.
        assert by_kind["publications"]["granularity"] == "year"
        assert by_kind["teaching-courses"]["granularity"] == "year"
        assert by_kind["research-contracts"]["granularity"] == "date"

    @pytest.mark.parametrize("kind", ALL_KINDS)
    def test_moi_muc_xuat_duoc_file_hop_le(self, client, as_role, kind):
        as_role("leader")
        res = client.get(f"{_BASE}/{kind}.xlsx")
        assert res.status_code == 200, f"{kind}: {res.text[:200]}"
        rows = _sheet(res)
        assert rows[0][0], "sheet phải có tiêu đề"
        assert _meta(rows, "Khoảng thời gian") == "Toàn bộ"

    def test_muc_khong_ton_tai_bao_404(self, client, as_role):
        as_role("leader")
        res = client.get(f"{_BASE}/khong-co-that.xlsx")
        assert res.status_code == 404, res.text


@requires_db
class TestLocTheoKhoangThoiGian:
    def _hop_dong(self, client, title: str, signed: str):
        res = client.post("/api/v1/research-contracts", json={
            "title": title, "contract_type": "nckh", "signed_date": signed,
        })
        assert res.status_code == 201, res.text
        return res.json()["data"]

    def test_chi_lay_ban_ghi_trong_khoang(self, client, as_role):
        """Hợp đồng lọc theo NGÀY KÝ."""
        as_role("admin")
        self._hop_dong(client, "HĐ trong khoảng", "2024-06-15")
        self._hop_dong(client, "HĐ trước khoảng", "2023-01-10")
        self._hop_dong(client, "HĐ sau khoảng", "2025-11-20")

        rows = _sheet(client.get(
            f"{_BASE}/research-contracts.xlsx", params={"from": "2024-01-01", "to": "2024-12-31"}
        ))
        assert _meta(rows, "Số bản ghi") == 1
        titles = [r[0] for r in _data_rows(rows) if r[0]]
        assert titles == ["HĐ trong khoảng"]

    def test_bien_khoang_la_bao_gom(self, client, as_role):
        """Đúng ngày đầu và ngày cuối phải nằm TRONG khoảng."""
        as_role("admin")
        self._hop_dong(client, "Đúng ngày đầu", "2024-01-01")
        self._hop_dong(client, "Đúng ngày cuối", "2024-12-31")

        rows = _sheet(client.get(
            f"{_BASE}/research-contracts.xlsx", params={"from": "2024-01-01", "to": "2024-12-31"}
        ))
        assert _meta(rows, "Số bản ghi") == 2

    def test_chi_co_from_hoac_chi_co_to(self, client, as_role):
        as_role("admin")
        self._hop_dong(client, "Cũ", "2023-05-05")
        self._hop_dong(client, "Mới", "2025-05-05")

        chi_tu = _sheet(client.get(f"{_BASE}/research-contracts.xlsx", params={"from": "2024-01-01"}))
        assert _meta(chi_tu, "Số bản ghi") == 1
        chi_den = _sheet(client.get(f"{_BASE}/research-contracts.xlsx", params={"to": "2024-01-01"}))
        assert _meta(chi_den, "Số bản ghi") == 1

    def test_khoang_nguoc_bi_tu_choi(self, client, as_role):
        as_role("admin")
        res = client.get(
            f"{_BASE}/research-contracts.xlsx", params={"from": "2025-01-01", "to": "2024-01-01"}
        )
        assert res.status_code == 400, res.text

    def test_ban_ghi_thieu_ngay_bi_loai_NHUNG_duoc_dem(self, client, as_role):
        """Bỏ bớt dòng mà không nói là cách nhanh nhất làm hỏng một báo cáo."""
        as_role("admin")
        self._hop_dong(client, "Có ngày ký", "2024-03-03")
        res = client.post("/api/v1/research-contracts", json={
            "title": "Chưa có ngày ký", "contract_type": "nckh",
        })
        assert res.status_code == 201, res.text

        rows = _sheet(client.get(
            f"{_BASE}/research-contracts.xlsx", params={"from": "2024-01-01", "to": "2024-12-31"}
        ))
        assert _meta(rows, "Số bản ghi") == 1
        assert _meta(rows, "Bị loại (không có dữ liệu thời gian)") == 1

    def test_khong_loc_thi_lay_ca_ban_ghi_thieu_ngay(self, client, as_role):
        as_role("admin")
        client.post("/api/v1/research-contracts", json={
            "title": "Không ngày", "contract_type": "nckh",
        })
        rows = _sheet(client.get(f"{_BASE}/research-contracts.xlsx"))
        assert _meta(rows, "Số bản ghi") == 1
        assert _meta(rows, "Bị loại (không có dữ liệu thời gian)") is None


@requires_db
class TestMucChiLuuNam:
    """publications / teaching-courses / student-mentorships không có ngày, chỉ có năm."""

    def test_khoang_ngay_quy_ve_khoang_nam(self, client, as_role, department):
        as_role("admin")
        for title, year in (("Bài 2023", 2023), ("Bài 2024", 2024), ("Bài 2025", 2025)):
            res = client.post("/api/v1/publications", json={
                "type": "conference", "title": title, "year": year,
                "journal": "Kỷ yếu Hội nghị KHCN",
                "authors": [{"external_name": "Nguyễn Văn A", "author_order": 1}],
            })
            assert res.status_code == 201, res.text

        # 01/06/2024 → 31/03/2025 phủ hai năm 2024 và 2025.
        rows = _sheet(client.get(
            f"{_BASE}/publications.xlsx", params={"from": "2024-06-01", "to": "2025-03-31"}
        ))
        assert _meta(rows, "Số bản ghi") == 2

    def test_sheet_noi_ro_da_quy_ve_nam(self, client, as_role):
        """Không nói thì người đọc tưởng đã lọc chính xác theo ngày."""
        as_role("admin")
        rows = _sheet(client.get(
            f"{_BASE}/publications.xlsx", params={"from": "2024-06-01", "to": "2025-03-31"}
        ))
        tieu_chi = _meta(rows, "Tiêu chí lọc")
        assert "Năm công bố" in tieu_chi
        assert "quy về khoảng năm" in tieu_chi


@requires_db
class TestPhamViXemQuyetDinhNoiDungFile:
    def test_ktv_chi_xuat_duoc_ban_ghi_cua_minh(self, client, as_role, department):
        """File = đúng thứ người dùng thấy trên màn hình, không hơn."""
        as_role("admin")
        res = client.post("/api/v1/publications", json={
            "type": "conference", "title": "Bài của người khác", "year": 2024,
            "journal": "Kỷ yếu Hội nghị KHCN",
            "authors": [{"external_name": "Người Ngoài", "author_order": 1}],
        })
        assert res.status_code == 201, res.text

        as_role("staff", department_id=department.id)
        rows = _sheet(client.get(f"{_BASE}/publications.xlsx"))
        assert _meta(rows, "Số bản ghi") == 0, "KTV xuất được bài mình không phải tác giả"

    def test_ten_file_kem_khoang_thoi_gian(self, client, as_role):
        as_role("leader")
        res = client.get(
            f"{_BASE}/community-services.xlsx", params={"from": "2024-01-01", "to": "2024-12-31"}
        )
        cd = res.headers["content-disposition"]
        assert "phuc-vu-cong-dong" in cd
        assert "2024-01-01" in cd and "2024-12-31" in cd


@requires_db
class TestVetKiemToan:
    def test_moi_lan_tai_deu_de_lai_vet(self, client, as_role, audit_rows):
        as_role("leader")
        truoc = audit_rows("RESEARCH_LIST_EXPORT")
        client.get(f"{_BASE}/community-services.xlsx")
        assert audit_rows("RESEARCH_LIST_EXPORT") == truoc + 1

    def test_ngay_thang_va_so_dong_vao_nhat_ky(self, client, as_role, db):
        from sqlalchemy import text as sa_text

        as_role("leader")
        client.get(
            f"{_BASE}/research-contracts.xlsx", params={"from": "2024-01-01", "to": "2024-12-31"}
        )
        detail = db.execute(sa_text(
            "SELECT detail FROM audit_logs WHERE action='RESEARCH_LIST_EXPORT' "
            "ORDER BY at DESC LIMIT 1"
        )).scalar_one()
        assert detail["kind"] == "research-contracts"
        assert detail["from"] == "2024-01-01" and detail["to"] == "2024-12-31"
        assert "rows" in detail

    def test_ngay_mai_khong_lam_vo_gi(self, client, as_role):
        """Khoảng tương lai là hợp lệ, chỉ ra 0 dòng."""
        as_role("leader")
        mai = (date.today() + timedelta(days=1)).isoformat()
        res = client.get(f"{_BASE}/research-contracts.xlsx", params={"from": mai})
        assert res.status_code == 200
        assert _meta(_sheet(res), "Số bản ghi") == 0
