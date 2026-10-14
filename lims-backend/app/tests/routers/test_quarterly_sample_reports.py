"""Báo cáo mẫu phân tích theo quý (m54) — luồng lab nhập → nộp → Văn phòng trả lại / chốt.

BẤT BIẾN
1. Lab chỉ sửa được báo cáo khi nó ở tay mình (draft / returned); đã nộp hay đã chốt
   thì mọi thao tác ghi dòng mẫu đều bị từ chối.
2. Ngày nhận mẫu phải nằm trong quý của báo cáo — quý tính theo ngày nhận.
3. Phòng này không xem, không sửa được báo cáo của phòng khác.
4. Văn phòng không sửa dữ liệu: chỉ trả lại (có lý do), chốt, mở khoá (có lý do).
5. File Excel mặc định chỉ gồm báo cáo đã chốt, và nói rõ phòng nào chưa có.
"""
import io
import uuid
from datetime import timedelta

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.models.department import Department
from app.models.notification import Notification
from app.models.sample_flow import TestParameter as CatalogParameter
from app.services.quarterly_sample_report_service import quarter_bounds, today_vn
from app.tests.conftest import requires_db

_BASE = "/api/v1/quarterly-sample-reports"


@pytest.fixture(autouse=True)
def _reset_rate_limit():
    from app.core.redis_client import get_redis

    r = get_redis()
    for k in r.scan_iter("ratelimit:quarterly-sample-export:*"):
        r.delete(k)
    yield


def _period():
    t = today_vn()
    q = (t.month - 1) // 3 + 1
    start, _ = quarter_bounds(t.year, q)
    return t.year, q, start


def _lab(db, name: str) -> Department:
    d = Department(name=name, code=f"T-{uuid.uuid4().hex[:6]}", kind="lab")
    db.add(d)
    db.flush()
    return d


def _param(db, matrix: str, name: str) -> CatalogParameter:
    p = CatalogParameter(matrix=matrix, name=name)
    db.add(p)
    db.flush()
    return p


def _item(**over):
    _, _, start = _period()
    body = {
        "sample_code": "26N265", "received_date": start.isoformat(), "matrix": "soil",
        "quantity": 2, "test_parameter_ids": [], "extra_criteria": "Tro",
    }
    body.update(over)
    return body


@requires_db
def test_full_workflow(client, db, as_role):
    year, quarter, start = _period()
    lab = _lab(db, "Hóa sinh – Hoạt chất sinh học ứng dụng (test)")
    moisture = _param(db, "soil", "Độ ẩm")
    water_ph = _param(db, "water", "pH")

    # Tài khoản Văn phòng có sẵn trước khi lab nộp — để nhận được thông báo.
    office = as_role("office")

    # ── Lab lập báo cáo và nhập mẫu ──
    as_role("staff", department_id=lab.id)
    rep = client.post(_BASE, json={"year": year, "quarter": quarter}).json()["data"]
    assert rep["status"] == "draft" and rep["can_edit"]
    again = client.post(_BASE, json={"year": year, "quarter": quarter}).json()["data"]
    assert again["id"] == rep["id"]  # gọi lại không lập bản thứ hai
    rid = rep["id"]

    res = client.post(f"{_BASE}/{rid}/items", json=_item(test_parameter_ids=[str(moisture.id)]))
    assert res.status_code == 201, res.text
    item = res.json()["data"]
    assert item["criteria_text"] == "Độ ẩm, Tro"
    assert item["matrix_label"] == "Đất"

    # Ngày ngoài quý, chỉ tiêu sai loại mẫu, không có chỉ tiêu → 422.
    out_of_q = (start - timedelta(days=1)).isoformat()
    assert client.post(f"{_BASE}/{rid}/items", json=_item(received_date=out_of_q)).status_code == 422
    assert client.post(
        f"{_BASE}/{rid}/items", json=_item(test_parameter_ids=[str(water_ph.id)])
    ).status_code == 422
    assert client.post(f"{_BASE}/{rid}/items", json=_item(extra_criteria=None)).status_code == 422
    assert client.post(f"{_BASE}/{rid}/items", json=_item(matrix="other")).status_code == 422

    # Cùng mã mẫu nhập lần hai vẫn được — một mẫu có thể nhiều đợt / nhiều phòng.
    assert client.post(f"{_BASE}/{rid}/items", json=_item(quantity=1)).status_code == 201

    # ── Nộp ──
    res = client.post(f"{_BASE}/{rid}/submit", json={"lab_note": "Đủ số liệu"})
    assert res.json()["data"]["status"] == "submitted"
    assert not res.json()["data"]["can_edit"]
    assert client.post(f"{_BASE}/{rid}/items", json=_item()).status_code == 409
    notified = db.execute(
        select(func.count()).select_from(Notification).where(
            Notification.user_id == office.id, Notification.ref_id == uuid.UUID(rid),
        )
    ).scalar_one()
    assert notified == 1

    # ── Văn phòng xem tổng quan rồi trả lại ──
    as_role("office")
    rows = client.get(f"{_BASE}/overview", params={"year": year, "quarter": quarter}).json()["data"]
    mine = next(r for r in rows if r["department"]["id"] == str(lab.id))
    assert (mine["status"], mine["item_count"], mine["total_quantity"]) == ("submitted", 2, 3)
    assert client.post(f"{_BASE}/{rid}/return", json={"reason": ""}).status_code == 400  # lỗi schema → 400
    res = client.post(f"{_BASE}/{rid}/return", json={"reason": "Mẫu 26N265 thiếu chỉ tiêu"})
    assert res.json()["data"]["status"] == "returned"
    assert res.json()["data"]["return_reason"] == "Mẫu 26N265 thiếu chỉ tiêu"
    assert client.post(f"{_BASE}/{rid}/items", json=_item()).status_code == 403  # VP không sửa

    # ── Lab (trưởng lab) sửa rồi nộp lại, Văn phòng chốt ──
    as_role("lab_manager", department_id=lab.id)
    items = client.get(f"{_BASE}/{rid}").json()["data"]["items"]
    upd = client.put(f"{_BASE}/{rid}/items/{items[0]['id']}", json=_item(extra_criteria="Tro, Đạm"))
    assert upd.status_code == 200, upd.text
    assert client.post(f"{_BASE}/{rid}/submit", json={}).json()["data"]["status"] == "submitted"

    as_role("office")
    assert client.post(f"{_BASE}/{rid}/lock").json()["data"]["status"] == "locked"
    assert client.post(f"{_BASE}/{rid}/lock").status_code == 409  # chốt hai lần

    # Đã chốt: lab không sửa được…
    as_role("staff", department_id=lab.id)
    assert client.delete(f"{_BASE}/{rid}/items/{items[0]['id']}").status_code == 409

    # …cho tới khi Văn phòng mở khoá (có lý do) — báo cáo về lại tay lab.
    as_role("office")
    res = client.post(f"{_BASE}/{rid}/reopen", json={"reason": "Sai số lượng mẫu 26N265"})
    assert res.json()["data"]["status"] == "returned"
    assert res.json()["data"]["locked_at"] is None
    as_role("staff", department_id=lab.id)
    assert client.delete(f"{_BASE}/{rid}/items/{items[0]['id']}").status_code == 204


@requires_db
def test_submit_empty_requires_confirmation(client, db, as_role):
    year, quarter, _ = _period()
    lab = _lab(db, "Lab rỗng (test)")
    as_role("staff", department_id=lab.id)
    rid = client.post(_BASE, json={"year": year, "quarter": quarter}).json()["data"]["id"]
    assert client.post(f"{_BASE}/{rid}/submit", json={}).status_code == 422
    res = client.post(f"{_BASE}/{rid}/submit", json={"confirm_empty": True})
    assert res.json()["data"]["status"] == "submitted"


@requires_db
def test_other_lab_cannot_read_or_write(client, db, as_role):
    year, quarter, _ = _period()
    lab_a, lab_b = _lab(db, "Lab A (test)"), _lab(db, "Lab B (test)")
    as_role("staff", department_id=lab_a.id)
    rid = client.post(_BASE, json={"year": year, "quarter": quarter}).json()["data"]["id"]

    as_role("staff", department_id=lab_b.id)
    assert client.get(f"{_BASE}/{rid}").status_code == 403
    assert client.post(f"{_BASE}/{rid}/items", json=_item()).status_code == 403
    assert client.post(f"{_BASE}/{rid}/submit", json={"confirm_empty": True}).status_code == 403
    assert client.get(
        f"{_BASE}/current", params={"year": year, "quarter": quarter, "department_id": str(lab_a.id)}
    ).status_code == 403
    # Lab không trả lại / chốt được, cũng không xem tổng quan toàn Viện.
    assert client.post(f"{_BASE}/{rid}/lock").status_code == 403
    assert client.get(
        f"{_BASE}/overview", params={"year": year, "quarter": quarter}
    ).status_code == 403


@requires_db
def test_future_quarter_rejected(client, db, as_role):
    year, quarter, _ = _period()
    lab = _lab(db, "Lab tương lai (test)")
    as_role("staff", department_id=lab.id)
    nxt = (year + 1, 1) if quarter == 4 else (year, quarter + 1)
    res = client.post(_BASE, json={"year": nxt[0], "quarter": nxt[1]})
    assert res.status_code == 422


@requires_db
def test_export_only_locked_by_default(client, db, as_role):
    year, quarter, _ = _period()
    locked_lab, open_lab = _lab(db, "AA Lab đã chốt (test)"), _lab(db, "AB Lab chưa chốt (test)")
    for lab, code in ((locked_lab, "26N001"), (open_lab, "26N002")):
        as_role("staff", department_id=lab.id)
        rid = client.post(_BASE, json={"year": year, "quarter": quarter}).json()["data"]["id"]
        client.post(f"{_BASE}/{rid}/items", json=_item(sample_code=code))
        client.post(f"{_BASE}/{rid}/submit", json={})
        if lab is locked_lab:
            as_role("office")
            client.post(f"{_BASE}/{rid}/lock")

    as_role("office")
    res = client.get(f"{_BASE}/export.xlsx", params={"year": year, "quarter": quarter})
    assert res.status_code == 200, res.text
    wb = load_workbook(io.BytesIO(res.content))
    ws = wb.worksheets[0]
    assert [c.value for c in ws[1]][1:] == [
        "STT", "MÃ SỐ MẪU", "LOẠI MẪU", "SỐ LƯỢNG", "CHỈ TIÊU PHÂN TÍCH",
        "PHÒNG NGHIÊN CỨU PHỤ TRÁCH",
    ]
    codes = [ws.cell(row=r, column=3).value for r in range(2, ws.max_row + 1)]
    assert "26N001" in codes and "26N002" not in codes
    status = {
        row[1]: (row[2], row[7])
        for row in wb["Tình trạng nộp"].iter_rows(min_row=4, values_only=True)
    }
    assert status["AA Lab đã chốt (test)"] == ("Đã chốt", "Có")
    assert status["AB Lab chưa chốt (test)"] == ("Đã nộp", "Không")

    res = client.get(
        f"{_BASE}/export.xlsx", params={"year": year, "quarter": quarter, "include_unlocked": True}
    )
    ws = load_workbook(io.BytesIO(res.content)).worksheets[0]
    codes = [ws.cell(row=r, column=3).value for r in range(2, ws.max_row + 1)]
    assert "26N002" in codes

    # Lab không xuất được file toàn Viện.
    as_role("staff", department_id=open_lab.id)
    assert client.get(
        f"{_BASE}/export.xlsx", params={"year": year, "quarter": quarter}
    ).status_code == 403


@requires_db
def test_batch_all_or_nothing(client, db, as_role):
    year, quarter, start = _period()
    lab = _lab(db, "Lab nhập lô (test)")
    as_role("staff", department_id=lab.id)
    rid = client.post(_BASE, json={"year": year, "quarter": quarter}).json()["data"]["id"]
    out_of_q = (start - timedelta(days=1)).isoformat()

    # Dòng 2 sai ngày → không lưu dòng nào, lỗi chỉ đúng dòng 2.
    bad = [_item(sample_code="26N301"), _item(sample_code="26N302", received_date=out_of_q),
           _item(sample_code="26N303")]
    res = client.post(f"{_BASE}/{rid}/items/batch", json={"items": bad})
    assert res.status_code == 422, res.text
    details = res.json()["error"]["details"]
    assert [(d["row"], d["sample_code"]) for d in details] == [(2, "26N302")]
    assert client.get(f"{_BASE}/{rid}").json()["data"]["item_count"] == 0

    good = [_item(sample_code=f"26N30{i}", quantity=i) for i in range(1, 4)]
    res = client.post(f"{_BASE}/{rid}/items/batch", json={"items": good})
    assert res.status_code == 201, res.text
    data = res.json()["data"]
    assert (data["item_count"], data["total_quantity"]) == (3, 6)


@requires_db
def test_stats_counts_locked_only_by_default(client, db, as_role):
    year, quarter, _ = _period()
    locked_lab, open_lab = _lab(db, "Lab thống kê chốt (test)"), _lab(db, "Lab thống kê nộp (test)")
    as_role("staff", department_id=locked_lab.id)
    rid = client.post(_BASE, json={"year": year, "quarter": quarter}).json()["data"]["id"]
    client.post(f"{_BASE}/{rid}/items/batch", json={"items": [
        _item(sample_code="26N401", quantity=2, extra_criteria="Tro"),
        _item(sample_code="26N402", matrix="water", quantity=3, extra_criteria="tro, COD"),
    ]})
    client.post(f"{_BASE}/{rid}/submit", json={})
    as_role("office")
    client.post(f"{_BASE}/{rid}/lock")

    as_role("staff", department_id=open_lab.id)
    rid2 = client.post(_BASE, json={"year": year, "quarter": quarter}).json()["data"]["id"]
    client.post(f"{_BASE}/{rid2}/items", json=_item(sample_code="26N403", quantity=7))
    client.post(f"{_BASE}/{rid2}/submit", json={})
    # Lab không xem thống kê toàn Viện.
    assert client.get(f"{_BASE}/stats", params={"year": year}).status_code == 403

    as_role("office")
    res = client.get(f"{_BASE}/stats", params={"year": year, "quarter": quarter})
    assert res.status_code == 200, res.text
    d = res.json()["data"]
    mine = [c for c in d["cells"] if c["department_id"] in (str(locked_lab.id), str(open_lab.id))]
    assert sorted((c["matrix"], c["quantity"]) for c in mine) == [("soil", 2), ("water", 3)]
    tro = next(t for t in d["top_criteria"] if t["name"].casefold() == "tro")
    assert tro["items"] >= 2  # "Tro" và "tro" là một chỉ tiêu

    d = client.get(
        f"{_BASE}/stats", params={"year": year, "quarter": quarter, "include_unlocked": True}
    ).json()["data"]
    mine = [c for c in d["cells"] if c["department_id"] == str(open_lab.id)]
    assert [c["quantity"] for c in mine] == [7]
