"""Nhập / xuất Excel Thẻ vào PTN — khoá bố cục theo file gốc của Văn phòng.

File gốc ("DANH SÁCH SINH VIÊN ĐÃ ĐƯỢC CẤP THẺ VÀO PTN QUÝ I.2026") có sheet TỔNG
với tiêu đề gộp A1:J1, hàng tiêu đề ở dòng 3, và nhiều sheet con theo GVHD. Test
dựng workbook giống hệt bằng tay — không cần DB, không phụ thuộc file thật.
"""
import io
from datetime import date
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from app.core.exceptions import AppException
from app.services import lab_access_excel_service as svc

_HEADERS = [
    "STT", "HỌ", "TÊN", "LỚP", "MSSV", "EMAIL", "ĐĂNG KÝ SỬ DỤNG PTN",
    "MỤC ĐÍCH SỬ DỤNG", "GIÁO VIÊN HƯỚNG DẪN", "THỜI GIAN", "GHI CHÚ",
]


def _source_file() -> bytes:
    wb = Workbook()
    total = wb.active
    total.title = "QUÝ I.2026"
    total["A1"] = "DANH SÁCH DÃ ĐƯỢC CẤP THẺ VÀO PTN QUÝ I.2026 (TỔNG)"
    total.append([])
    total.append(_HEADERS)
    total.append([1, "Phan Minh ", "Thịnh", "DH23SHA ", 23126202, "23126202@@st.hcmuaf.edu.vn",
                  "RIBE  306", "Rèn Nghề ", "Nguyễn Thị Quyên ", "06/4/26-06/7/26", None])
    total.append([2, "Trịnh Thanh ", "Tình", "DH22SHA", 22126183, "22126183@st.hcmuaf.edu.vn",
                  "RIBE 316,318,320-A1; RIBE 304,306-A2", "KLTN", "TS. Trịnh Thị Phi Ly",
                  "24.3.2026-24.9.2026", "GH"])
    total.append([3, "Phạm Văn ", "Hiếu", "DH22SHA", 22126060, None, None, "Rèn Nghề ",
                  "TS. Phạm Đức Toàn", "06/4/26-06/7/26", None])
    total.append([4, "Lỗi", "Ngày", None, 1, None, None, None, None, "không rõ", None])
    total.append([5, "Trùng", "Mã", None, 23126202, None, None, None, None, "06/4/26", None])
    # Sheet con theo GVHD: trùng dữ liệu, không có THỜI GIAN — phải bị bỏ qua.
    sub = wb.create_sheet("Nguyễn Thị Quyên")
    sub.append(_HEADERS[:9])
    sub.append([1, "Phan Minh ", "Thịnh", "DH23SHA", 23126202, "x@y", "BIO 203", "Học nghề", "Q"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parse_reads_total_sheet_and_cleans_values():
    parsed = svc.parse_workbook(_source_file())
    assert parsed.sheet == "QUÝ I.2026"
    assert [r.data["student_code"] for r in parsed.rows] == ["23126202", "22126183", "22126060"]
    first = parsed.rows[0].data
    assert first["student_name"] == "Phan Minh Thịnh"
    assert first["email"] == "23126202@st.hcmuaf.edu.vn"  # "@@" gõ nhầm
    assert first["room"] == "RIBE 306"
    assert (first["valid_from"], first["valid_to"]) == (date(2026, 4, 6), date(2026, 7, 6))
    assert parsed.rows[1].data["note"] == "GH"
    assert parsed.rows[2].data["room"] is None  # file gốc có dòng chưa ghi phòng
    assert [(e["row"], e["message"][:9]) for e in parsed.errors] == [
        (7, "THỜI GIAN"), (8, "MSSV trùn"),
    ]


@pytest.mark.parametrize("raw,expected", [
    ("06/4/26-06/7/26", (date(2026, 4, 6), date(2026, 7, 6))),
    ("24.3.2026-24.9.2026", (date(2026, 3, 24), date(2026, 9, 24))),
    ("24/3/26 - 24/6/26", (date(2026, 3, 24), date(2026, 6, 24))),
    ("06/4/2026", (date(2026, 4, 6), None)),
    (date(2026, 1, 2), (date(2026, 1, 2), None)),
])
def test_parse_period(raw, expected):
    assert svc.parse_period(raw) == expected


@pytest.mark.parametrize("raw", [None, "", "abc", "31/2/26", "06/7/26-06/4/26"])
def test_parse_period_rejects(raw):
    with pytest.raises(ValueError):
        svc.parse_period(raw)


def test_parse_rejects_file_without_table():
    wb = Workbook()
    wb.active.append(["a", "b"])
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(AppException):
        svc.parse_workbook(buf.getvalue())
    with pytest.raises(AppException):
        svc.parse_workbook(b"not an xlsx")


def test_supervisor_display_strips_titles():
    assert svc.supervisor_display("TS.BSTY. Đinh Xuân Phát ") == "Đinh Xuân Phát"
    assert svc.supervisor_display("PGS. TS. Nguyễn Ngọc Tấn") == "Nguyễn Ngọc Tấn"
    assert svc.supervisor_display("Ths. Nguyễn Thị Quyên") == "Nguyễn Thị Quyên"
    assert svc.supervisor_display("Thạch Văn A") == "Thạch Văn A"


def test_export_layout_matches_source_and_round_trips():
    parsed = svc.parse_workbook(_source_file())
    cards = [SimpleNamespace(**r.data) for r in parsed.rows]
    wb = load_workbook(io.BytesIO(svc.build_workbook(cards, quarter=1, year=2026)))

    assert wb.sheetnames == ["QUÝ I.2026", "Nguyễn Thị Quyên", "Trịnh Thị Phi Ly", "Phạm Đức Toàn"]
    ws = wb.worksheets[0]
    assert ws["A1"].value == "DANH SÁCH ĐÃ ĐƯỢC CẤP THẺ VÀO PTN QUÝ I.2026 (TỔNG)"
    assert [str(r) for r in ws.merged_cells.ranges] == ["A1:J1"]
    assert [c.value for c in ws[3]] == _HEADERS
    assert [c.value for c in ws[4]] == [
        1, "Phan Minh", "Thịnh", "DH23SHA", 23126202, "23126202@st.hcmuaf.edu.vn",
        "RIBE 306", "Rèn Nghề", "Nguyễn Thị Quyên", "06/4/26-06/7/26", None,
    ]
    assert ws["J5"].value == "24/3/26-24/9/26"
    assert ws["A4"].font.name == "Times New Roman" and ws["B3"].font.b
    assert ws["A4"].border.left.style == "thin"
    assert ws["H5"].font.color.rgb == "FFFF0000"  # KLTN chữ đỏ như file gốc
    assert ws.row_dimensions[5].height == 28.0  # phòng dài → 2 dòng
    assert round(ws.column_dimensions["F"].width, 2) == 30.18

    # Nhập lại chính file vừa xuất phải ra đúng dữ liệu ban đầu.
    again = svc.parse_workbook(svc.build_workbook(cards, quarter=1, year=2026))
    assert [r.data for r in again.rows] == [r.data for r in parsed.rows]


def test_export_neutralises_formulas():
    card = SimpleNamespace(
        student_name="=HYPERLINK(1)", class_name=None, student_code="1", email=None,
        room=None, purpose=None, supervisor_name=None,
        valid_from=date(2026, 1, 1), valid_to=None, note=None,
    )
    ws = load_workbook(io.BytesIO(svc.build_workbook([card], quarter=None, year=None))).active
    assert ws.title == "TỔNG"
    assert ws["C4"].data_type == "s"
