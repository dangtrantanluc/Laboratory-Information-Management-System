"""Nhập / xuất Excel Thẻ vào PTN — lấy file của Văn phòng làm gốc.

File gốc: "DANH SÁCH SINH VIÊN ĐÃ ĐƯỢC CẤP THẺ VÀO PTN QUÝ I.2026". Một workbook gồm:
  - sheet TỔNG (vd "QUÝ I.2026"): tiêu đề gộp A1:J1, hàng tiêu đề cột ở dòng 3,
    11 cột STT · HỌ · TÊN · LỚP · MSSV · EMAIL · ĐĂNG KÝ SỬ DỤNG PTN · MỤC ĐÍCH SỬ DỤNG ·
    GIÁO VIÊN HƯỚNG DẪN · THỜI GIAN · GHI CHÚ;
  - mỗi giáo viên hướng dẫn một sheet con — chỉ là bản trích của sheet TỔNG.

NHẬP
Chỉ đọc sheet TỔNG (sheet con trùng dữ liệu, lại hay thiếu cột THỜI GIAN). Khoá đối
chiếu là MSSV: có rồi thì cập nhật theo file, chưa có thì tạo. Không xoá bản ghi
vắng mặt trong file — một file quý không phải toàn bộ lịch sử cấp thẻ.
Luôn chạy thử (dry_run) trước để người dùng xem sẽ tạo / sửa bao nhiêu dòng, dòng
nào lỗi, rồi mới ghi.

XUẤT
Dựng lại đúng bố cục đó: cùng font Times New Roman, độ rộng cột, chiều cao dòng,
viền mảnh, "KLTN" chữ đỏ, THỜI GIAN dạng "06/4/26-06/7/26"; kèm sheet con theo GVHD.
Thứ tự dòng = thứ tự tạo, nên nhập rồi xuất lại giữ nguyên thứ tự của file gốc.
"""
import io
import math
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta, timezone
from typing import Optional

from sqlalchemy import or_, select
from sqlalchemy.orm import Session

from app.core.concurrency import export_slot
from app.core.deps import CurrentUser
from app.core.error_codes import ErrorCode
from app.core.exceptions import AppException
from app.models.lab_access import LabAccessCard
from app.services import audit_service

IMPORT_MAX_BYTES = 5 * 1024 * 1024
IMPORT_MAX_ROWS = 5000
EXPORT_MAX_ROWS = 10000
_HEADER_SCAN_ROWS = 10

# ── Bố cục file gốc ───────────────────────────────────────────────
# (tiêu đề cột, độ rộng) — số đo lấy nguyên từ sheet TỔNG của file Quý I.2026.
_COLUMNS: list[tuple[str, float]] = [
    ("STT", 8.81640625),
    ("HỌ", 20.26953125),
    ("TÊN", 8.7265625),
    ("LỚP", 13.0),
    ("MSSV", 10.54296875),
    ("EMAIL", 30.1796875),
    ("ĐĂNG KÝ SỬ DỤNG PTN", 22.54296875),
    ("MỤC ĐÍCH SỬ DỤNG", 15.453125),
    ("GIÁO VIÊN HƯỚNG DẪN", 30.6328125),
    ("THỜI GIAN", 19.81640625),
    ("GHI CHÚ", 8.7265625),
]
_TITLE_ROW = 1
_HEADER_ROW = 3
_TITLE_HEIGHT = 17.5
_HEADER_HEIGHT = 30.0
_ROW_HEIGHT = 15.5
_LINE_HEIGHT = 14.0
_ROOM_CHARS_PER_LINE = 22  # cột G rộng ~22.5 ký tự, là cột duy nhất xuống dòng
_ZOOM = 85
_FONT = "Times New Roman"
_ROMAN = {1: "I", 2: "II", 3: "III", 4: "IV"}

# Tiêu đề cột (đã chuẩn hoá) → trường. Nhận thêm vài biến thể hay gặp.
_HEADER_ALIASES: dict[str, str] = {
    "STT": "stt",
    "HỌ": "last_name",
    "HỌ ĐỆM": "last_name",
    "HỌ VÀ TÊN LÓT": "last_name",
    "TÊN": "first_name",
    "HỌ VÀ TÊN": "full_name",
    "HỌ TÊN": "full_name",
    "LỚP": "class_name",
    "MSSV": "student_code",
    "MÃ SỐ SINH VIÊN": "student_code",
    "EMAIL": "email",
    "ĐĂNG KÝ SỬ DỤNG PTN": "room",
    "PHÒNG ĐĂNG KÝ": "room",
    "MỤC ĐÍCH SỬ DỤNG": "purpose",
    "MỤC ĐÍCH": "purpose",
    "GIÁO VIÊN HƯỚNG DẪN": "supervisor_name",
    "GVHD": "supervisor_name",
    "THỜI GIAN": "period",
    "GHI CHÚ": "note",
}

_COMPARE_FIELDS = (
    "student_name", "class_name", "email", "room", "purpose",
    "supervisor_name", "valid_from", "valid_to", "note",
)
_MAX_LEN = {
    "student_name": 255, "class_name": 50, "student_code": 30, "email": 255,
    "room": 255, "purpose": 255, "supervisor_name": 255, "note": 1000,
}

_DATE_RE = re.compile(r"(\d{1,2})\s*[./-]\s*(\d{1,2})\s*[./-]\s*(\d{4}|\d{2})(?!\d)")
_ACADEMIC_PREFIX_RE = re.compile(
    r"^(?:(?:PGS|GS|TS|THS|BSTY|BS|CN|KS|NCS)(?:\s*\.\s*|\s+))+", re.IGNORECASE
)
_SHEET_NAME_BAD = re.compile(r"[\[\]:*?/\\]")


# ── Tiện ích chuỗi ────────────────────────────────────────────────
def _norm_header(v) -> str:
    if v is None:
        return ""
    s = unicodedata.normalize("NFC", str(v))
    return " ".join(s.split()).upper()


def _text(v) -> Optional[str]:
    """Ô → chuỗi đã gọn khoảng trắng; số nguyên kiểu float (22126145.0) → '22126145'."""
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = " ".join(unicodedata.normalize("NFC", str(v)).split())
    return s or None


def _email(v) -> Optional[str]:
    s = _text(v)
    if not s:
        return None
    # File gốc có nhiều dòng gõ nhầm "@@" và thừa khoảng trắng.
    return re.sub(r"@+", "@", s.replace(" ", ""))


def _parse_date(d: str, m: str, y: str) -> date:
    year = int(y)
    if year < 100:
        year += 2000
    return date(year, int(m), int(d))


def parse_period(v) -> tuple[date, Optional[date]]:
    """'06/4/26-06/7/26' · '24.3.2026-24.9.2026' · ô ngày của Excel → (từ, đến)."""
    if isinstance(v, datetime):
        return v.date(), None
    if isinstance(v, date):
        return v, None
    s = _text(v)
    if not s:
        raise ValueError("Thiếu THỜI GIAN")
    found = _DATE_RE.findall(s)
    if not found:
        raise ValueError(f"THỜI GIAN không đọc được: '{s}'")
    try:
        start = _parse_date(*found[0])
        end = _parse_date(*found[1]) if len(found) > 1 else None
    except ValueError:
        raise ValueError(f"THỜI GIAN có ngày không hợp lệ: '{s}'") from None
    if end and end < start:
        raise ValueError(f"THỜI GIAN: ngày kết thúc trước ngày bắt đầu ('{s}')")
    return start, end


def format_period(start: date, end: Optional[date]) -> str:
    def f(d: date) -> str:
        return f"{d.day:02d}/{d.month}/{d.year % 100:02d}"

    return f"{f(start)}-{f(end)}" if end else f(start)


def split_name(full: str) -> tuple[str, str]:
    """'Nguyễn Thị Ngọc Diểm' → ('Nguyễn Thị Ngọc', 'Diểm') — TÊN là từ cuối."""
    parts = full.split()
    if len(parts) <= 1:
        return "", full.strip()
    return " ".join(parts[:-1]), parts[-1]


def supervisor_display(name: str) -> str:
    """Bỏ học hàm/học vị: 'TS.BSTY. Đinh Xuân Phát' → 'Đinh Xuân Phát'."""
    return " ".join(_ACADEMIC_PREFIX_RE.sub("", name.strip()).split())


# ── Đọc file ──────────────────────────────────────────────────────
@dataclass
class ParsedRow:
    row: int
    data: dict


@dataclass
class ParsedSheet:
    sheet: str
    rows: list[ParsedRow] = field(default_factory=list)
    errors: list[dict] = field(default_factory=list)


def _find_header(rows: list[tuple]) -> Optional[tuple[int, dict[str, int]]]:
    for idx, row in enumerate(rows[:_HEADER_SCAN_ROWS]):
        cols: dict[str, int] = {}
        for ci, v in enumerate(row):
            key = _HEADER_ALIASES.get(_norm_header(v))
            if key and key not in cols:
                cols[key] = ci
        has_name = "full_name" in cols or "first_name" in cols
        if "student_code" in cols and has_name and "period" in cols:
            return idx, cols
    return None


def parse_workbook(content: bytes) -> ParsedSheet:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise AppException(
            ErrorCode.INVALID_FILE_TYPE, "Không đọc được tệp. Hãy chọn tệp Excel .xlsx", 400
        ) from None

    candidates = []
    try:
        for ws in wb.worksheets:
            rows = []
            for i, r in enumerate(ws.iter_rows(values_only=True)):
                if i > IMPORT_MAX_ROWS + _HEADER_SCAN_ROWS:
                    raise AppException(
                        ErrorCode.VALIDATION_ERROR,
                        f"Sheet '{ws.title}' vượt quá {IMPORT_MAX_ROWS} dòng", 422,
                    )
                rows.append(r)
            found = _find_header(rows)
            if not found:
                continue
            title = " ".join(_norm_header(v) for v in (rows[0] if rows else ()) if v)
            is_total = "TỔNG" in title or "TỔNG" in _norm_header(ws.title)
            candidates.append((is_total, ws.title, rows, found))
    finally:
        wb.close()

    if not candidates:
        raise AppException(
            ErrorCode.VALIDATION_ERROR,
            "Không tìm thấy bảng có đủ cột MSSV, HỌ/TÊN và THỜI GIAN trong tệp", 422,
        )
    # Ưu tiên sheet TỔNG; không có thì lấy sheet hợp lệ đầu tiên.
    _, sheet, rows, (header_idx, cols) = next(
        (c for c in candidates if c[0]), candidates[0]
    )
    out = ParsedSheet(sheet=sheet)
    seen: dict[str, int] = {}

    def cell(row: tuple, key: str):
        ci = cols.get(key)
        return row[ci] if ci is not None and ci < len(row) else None

    for offset, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        if not any(_text(v) for v in row):
            continue
        code = _text(cell(row, "student_code"))
        if "full_name" in cols:
            name = _text(cell(row, "full_name"))
        else:
            name = _text(" ".join(filter(None, (
                _text(cell(row, "last_name")), _text(cell(row, "first_name")),
            ))))

        def err(msg: str) -> None:
            out.errors.append({"row": offset, "student_code": code, "message": msg})

        if not code:
            err("Thiếu MSSV")
            continue
        if not name:
            err("Thiếu họ tên")
            continue
        if code in seen:
            err(f"MSSV trùng với dòng {seen[code]}")
            continue
        try:
            valid_from, valid_to = parse_period(cell(row, "period"))
        except ValueError as e:
            err(str(e))
            continue
        data = {
            "student_code": code,
            "student_name": name,
            "class_name": _text(cell(row, "class_name")),
            "email": _email(cell(row, "email")),
            "room": _text(cell(row, "room")),
            "purpose": _text(cell(row, "purpose")),
            "supervisor_name": _text(cell(row, "supervisor_name")),
            "valid_from": valid_from,
            "valid_to": valid_to,
            "note": _text(cell(row, "note")),
        }
        too_long = [k for k, n in _MAX_LEN.items() if data.get(k) and len(data[k]) > n]
        if too_long:
            err("Quá dài: " + ", ".join(too_long))
            continue
        seen[code] = offset
        out.rows.append(ParsedRow(row=offset, data=data))
    return out


# ── Nhập ──────────────────────────────────────────────────────────
def import_cards(
    db: Session,
    *,
    user: CurrentUser,
    content: bytes,
    dry_run: bool,
    correlation_id: Optional[str],
    ip: Optional[str],
) -> dict:
    if len(content) > IMPORT_MAX_BYTES:
        raise AppException(ErrorCode.FILE_TOO_LARGE, "Tệp vượt quá 5 MB", 413)
    parsed = parse_workbook(content)

    codes = [r.data["student_code"] for r in parsed.rows]
    existing: dict[str, LabAccessCard] = {}
    if codes:
        # Một MSSV có thể có nhiều thẻ (cấp lại) — đối chiếu với thẻ mới nhất.
        for c in db.execute(
            select(LabAccessCard)
            .where(LabAccessCard.student_code.in_(codes))
            .order_by(LabAccessCard.valid_from.desc(), LabAccessCard.created_at.desc())
        ).scalars():
            existing.setdefault(c.student_code, c)

    # created_at tăng dần theo thứ tự dòng: xuất lại giữ đúng thứ tự file gốc.
    base = datetime.now(timezone.utc)
    items: list[dict] = []
    counts = {"create": 0, "update": 0, "unchanged": 0}
    for i, r in enumerate(parsed.rows):
        card = existing.get(r.data["student_code"])
        if card is None:
            action = "create"
            if not dry_run:
                db.add(LabAccessCard(
                    **r.data, created_by=user.id, updated_by=user.id,
                    created_at=base + timedelta(microseconds=i),
                ))
        else:
            changed = {k: r.data[k] for k in _COMPARE_FIELDS if getattr(card, k) != r.data[k]}
            action = "update" if changed else "unchanged"
            if changed and not dry_run:
                for k, v in changed.items():
                    setattr(card, k, v)
                card.updated_by = user.id
        counts[action] += 1
        items.append({
            "row": r.row, "student_code": r.data["student_code"],
            "student_name": r.data["student_name"], "action": action,
        })

    result = {
        "dry_run": dry_run,
        "sheet": parsed.sheet,
        "total_rows": len(parsed.rows) + len(parsed.errors),
        "created": counts["create"],
        "updated": counts["update"],
        "unchanged": counts["unchanged"],
        "errors": parsed.errors,
        "items": items,
    }
    if not dry_run:
        audit_service.log_action(
            db, action="LAB_ACCESS_CARD_IMPORT", resource="lab_access_card",
            user_id=user.id, resource_id=None, correlation_id=correlation_id, ip=ip,
            detail={
                "sheet": parsed.sheet, "created": counts["create"],
                "updated": counts["update"], "errors": len(parsed.errors),
            },
        )
        db.commit()
    return result


# ── Xuất ──────────────────────────────────────────────────────────
def _sheet_name(raw: str, used: set[str]) -> str:
    base = _SHEET_NAME_BAD.sub(" ", raw).strip()[:31] or "Sheet"
    name, n = base, 2
    while name.casefold() in used:
        suffix = f" ({n})"
        name = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(name.casefold())
    return name


def _student_code_cell(code: str):
    # File gốc lưu MSSV dạng số; giữ chuỗi nếu có số 0 đầu để không mất ký tự.
    if code.isdigit() and not code.startswith("0") and len(code) <= 15:
        return int(code)
    return code


def _write_sheet(ws, cards: list[LabAccessCard], title: Optional[str]) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    white = PatternFill("solid", fgColor="FFFFFFFF")
    center = Alignment(horizontal="center")
    wrap_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    body = Font(name=_FONT, size=12)
    red = Font(name=_FONT, size=12, color="FFFF0000")
    bold = Font(name=_FONT, size=12, bold=True)

    for i, (_, width) in enumerate(_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    if title:
        c = ws.cell(row=_TITLE_ROW, column=1, value=title)
        c.font = Font(name=_FONT, size=14, bold=True)
        c.alignment = center
        c.fill = white
        ws.merge_cells(start_row=_TITLE_ROW, start_column=1, end_row=_TITLE_ROW, end_column=10)
        ws.row_dimensions[_TITLE_ROW].height = _TITLE_HEIGHT

    for ci, (label, _) in enumerate(_COLUMNS, start=1):
        c = ws.cell(row=_HEADER_ROW, column=ci, value=label)
        c.font, c.alignment, c.border, c.fill = bold, wrap_center, border, white
    ws.row_dimensions[_HEADER_ROW].height = _HEADER_HEIGHT

    for n, card in enumerate(cards, start=1):
        r = _HEADER_ROW + n
        ho, ten = split_name(card.student_name)
        values = [
            n, ho, ten, card.class_name, _student_code_cell(card.student_code), card.email,
            card.room, card.purpose, card.supervisor_name,
            format_period(card.valid_from, card.valid_to), card.note,
        ]
        for ci, v in enumerate(values, start=1):
            c = ws.cell(row=r, column=ci, value=v)
            if isinstance(v, str):
                # Ghi thuần văn bản: chuỗi bắt đầu bằng '=' không thành công thức.
                c.data_type = "s"
            c.border, c.fill = border, white
            c.alignment = wrap_center if ci == 7 else center
            c.font = red if ci == 8 and (v or "").strip().upper() == "KLTN" else body
        lines = max(1, math.ceil(len(card.room or "") / _ROOM_CHARS_PER_LINE))
        ws.row_dimensions[r].height = _ROW_HEIGHT if lines == 1 else _LINE_HEIGHT * lines


def build_workbook(
    cards: list[LabAccessCard], *, quarter: Optional[int], year: Optional[int]
) -> bytes:
    from openpyxl import Workbook

    period = f"QUÝ {_ROMAN[quarter]}.{year}" if quarter and year else None
    wb = Workbook()
    ws = wb.active
    used: set[str] = set()
    ws.title = _sheet_name(period or "TỔNG", used)
    ws.sheet_view.zoomScale = _ZOOM
    title = "DANH SÁCH ĐÃ ĐƯỢC CẤP THẺ VÀO PTN"
    _write_sheet(ws, cards, f"{title} {period} (TỔNG)" if period else f"{title} (TỔNG)")

    # Sheet con theo GVHD — thứ tự theo lần xuất hiện đầu tiên trong danh sách tổng.
    groups: dict[str, tuple[str, list[LabAccessCard]]] = {}
    for c in cards:
        if not c.supervisor_name:
            continue
        display = supervisor_display(c.supervisor_name)
        if not display:
            continue
        groups.setdefault(display.casefold(), (display, []))[1].append(c)
    for display, group in groups.values():
        _write_sheet(wb.create_sheet(_sheet_name(display, used)), group, None)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_xlsx(
    db: Session,
    *,
    user: CurrentUser,
    q: Optional[str],
    supervisor_name: Optional[str],
    quarter: Optional[int],
    year: Optional[int],
    correlation_id: Optional[str],
    ip: Optional[str],
) -> tuple[bytes, str]:
    with export_slot():
        conds = []
        if q:
            like = f"%{q.strip()}%"
            conds.append(or_(
                LabAccessCard.student_name.ilike(like),
                LabAccessCard.student_code.ilike(like),
                LabAccessCard.class_name.ilike(like),
                LabAccessCard.email.ilike(like),
            ))
        if supervisor_name:
            conds.append(LabAccessCard.supervisor_name.ilike(f"%{supervisor_name.strip()}%"))
        cards = db.execute(
            select(LabAccessCard)
            .where(*conds)
            .order_by(LabAccessCard.created_at, LabAccessCard.student_name)
            .limit(EXPORT_MAX_ROWS + 1)
        ).scalars().all()
        if len(cards) > EXPORT_MAX_ROWS:
            raise AppException(
                ErrorCode.EXPORT_TOO_LARGE,
                f"Vượt ngưỡng {EXPORT_MAX_ROWS} dòng. Thu hẹp bộ lọc.", 422,
            )
        content = build_workbook(list(cards), quarter=quarter, year=year)
        audit_service.log_action(
            db, action="LAB_ACCESS_CARD_EXPORT", resource="lab_access_card",
            user_id=user.id, resource_id=None, correlation_id=correlation_id, ip=ip,
            detail={"rows": len(cards), "quarter": quarter, "year": year},
        )
        db.commit()
    suffix = f"quy-{_ROMAN[quarter]}-{year}" if quarter and year else date.today().isoformat()
    return content, f"the-vao-ptn-{suffix}.xlsx"

