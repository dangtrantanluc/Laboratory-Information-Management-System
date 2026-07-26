"""M6 export service (M6.4) — xuất Excel/PDF báo cáo tổng hợp (#12/#13).

Áp RBAC scope vào NỘI DUNG file (NFR-SEC-RPT-002): office không xuất samples (403);
staff chỉ phòng mình + không cột tiền; system-access chỉ admin/leader. Ghi audit
REPORT_EXPORT (BR-RPT-012). Tái dùng service tổng hợp (dashboard/report/system-access)
để file = đúng dữ liệu trên màn (consistency).
"""
import io
import logging
from typing import Optional

from sqlalchemy.orm import Session

from app.core.deps import CurrentUser
from app.services import (
    audit_service,
    dashboard_service,
    unified_report_service,
    report_common as rc,
    system_access_service,
)

logger = logging.getLogger("lims.report_export")

SUPPORTED_REPORT_TYPES = ("dashboard", "samples", "chemicals", "system-access")
PDF_SUPPORTED = ("dashboard",)  # bản đầu chỉ dashboard có template PDF (OQ#3)


def _check_report_type(report_type: str) -> None:
    if report_type not in SUPPORTED_REPORT_TYPES:
        raise rc.err(
            "REPORT_TYPE_NOT_FOUND",
            f"Loại báo cáo '{report_type}' không hỗ trợ", 404,
        )


def _enforce_export_rbac(db: Session, user: CurrentUser, report_type: str) -> None:
    if report_type == "samples":
        rc.deny_office_samples(user)  # office 403
    if report_type == "system-access":
        rc.require_audit_read(db, user)  # chỉ admin/leader


def _gather(
    db: Session, user: CurrentUser, report_type: str, params: dict
) -> dict:
    """Lấy data đã áp scope từ service tổng hợp tương ứng."""
    if report_type == "dashboard":
        data, _meta = dashboard_service.get_dashboard(
            db, user=user,
            department_id=params.get("department_id"),
            due_within_days=params.get("due_within_days", 30),
        )
        charts, _meta2 = dashboard_service.get_charts(
            db, user=user,
            date_from=params.get("date_from"), date_to=params.get("date_to"),
            department_id=params.get("department_id"),
            group_by=params.get("group_by", "month"),
            charts=None,
        )
        data["charts"] = charts
        return data
    if report_type == "samples":
        data, _meta = unified_report_service.report_samples(
            db, user=user,
            date_from=params.get("date_from"), date_to=params.get("date_to"),
            department_id=params.get("department_id"),
            group_by=params.get("group_by", "month"),
            status_filter=params.get("status"),
            time_field=params.get("time_field", "received_at"),
            breakdown=params.get("breakdown", "status"),
        )
        return data
    if report_type == "chemicals":
        data, _meta = unified_report_service.report_chemicals(
            db, user=user,
            date_from=params.get("date_from"), date_to=params.get("date_to"),
            department_id=params.get("department_id"),
            group_by=params.get("group_by", "month"),
            measurement_group=params.get("type"),
            chemical_id=params.get("chemical_id"),
            metric=params.get("metric", "consumption"),
        )
        return data
    # system-access
    data, _meta = system_access_service.system_access(
        db, user=user,
        date_from=params.get("date_from"), date_to=params.get("date_to"),
        group_by=params.get("group_by", "month"),
        filter_user_id=params.get("user_id"),
        action_type=params.get("action_type", "all"),
        top_n=params.get("top_n", 10),
        include_timeline=True,
    )
    return data


def _audit_export(
    db: Session, user: CurrentUser, report_type: str, fmt: str,
    params: dict, correlation_id: Optional[str], ip: Optional[str],
) -> None:
    audit_service.log_action(
        db,
        action="REPORT_EXPORT",
        resource="report",
        user_id=user.id,
        correlation_id=correlation_id,
        ip=ip,
        detail={
            "report_type": report_type,
            "format": fmt,
            "from": str(params.get("date_from")) if params.get("date_from") else None,
            "to": str(params.get("date_to")) if params.get("date_to") else None,
            "scope_role": user.role,
            "scope_department": str(params.get("department_id"))
            if params.get("department_id") else None,
        },
    )
    db.commit()


# ======================= XLSX (#12) =======================
_FORMULA_TRIGGERS = ("=", "+", "-", "@")


def _safe_cell(v):
    """Chặn Excel/CSV formula injection: prefix nháy đơn nếu chuỗi bắt đầu bằng ký tự trigger công thức."""
    if isinstance(v, str) and v.startswith(_FORMULA_TRIGGERS):
        return "'" + v
    return v


def _append_row(ws, row: list) -> None:
    ws.append([_safe_cell(v) for v in row])


def export_xlsx(
    db: Session, *, user: CurrentUser, report_type: str, params: dict,
    correlation_id: Optional[str], ip: Optional[str],
) -> tuple[bytes, str]:
    _check_report_type(report_type)
    _enforce_export_rbac(db, user, report_type)
    data = _gather(db, user, report_type, params)

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = report_type[:31]
    _append_row(ws, [f"Báo cáo: {report_type}", f"Người xuất: {user.full_name}", f"Vai trò: {user.role}"])
    _write_sheet(ws, report_type, data)

    _audit_export(db, user, report_type, "xlsx", params, correlation_id, ip)
    buf = io.BytesIO()
    wb.save(buf)
    filename = f"bao-cao-{report_type}.xlsx"
    return buf.getvalue(), filename


def _write_sheet(ws, report_type: str, data: dict) -> None:
    _append_row(ws, [])
    if report_type == "dashboard":
        for block, content in data.items():
            if block == "scope" or not isinstance(content, dict):
                continue
            _append_row(ws, [f"== {block} =="])
            for k, v in content.items():
                if isinstance(v, (dict, list)):
                    _append_row(ws, [k, str(v)])
                else:
                    _append_row(ws, [k, v])
            _append_row(ws, [])
    elif report_type == "samples":
        _append_row(ws, ["Tổng số mẫu", data.get("total", 0)])
        _append_row(ws, [])
        _append_row(ws, ["Trạng thái", "Số lượng"])
        for s, c in (data.get("by_status") or {}).items():
            _append_row(ws, [s, c])
        _append_row(ws, [])
        _append_row(ws, ["Kỳ", "Số mẫu"])
        for row in data.get("by_time", []):
            _append_row(ws, [row["period"], row["count"]])
    elif report_type == "chemicals":
        has_cost = any("consumption_cost" in g for g in data.get("by_measurement_group", []))
        header = ["Nhóm đo", "Đơn vị base", "Tổng lượng"] + (["Chi phí"] if has_cost else [])
        _append_row(ws, header)
        for g in data.get("by_measurement_group", []):
            row = [g["measurement_group"], g["base_unit"], g["total_qty"]]
            if has_cost:
                row.append(g.get("consumption_cost", ""))
            _append_row(ws, row)
        if "total_cost" in data:
            _append_row(ws, ["Tổng chi phí", "", "", data["total_cost"]])
    else:  # system-access
        totals = data.get("totals", {})
        _append_row(ws, ["Lượt truy cập", totals.get("access_count", 0)])
        _append_row(ws, ["Lượt tải", totals.get("download_count", 0)])
        _append_row(ws, ["Lượt chỉnh sửa", totals.get("edit_count", 0)])
        _append_row(ws, [])
        _append_row(ws, ["Kỳ", "Truy cập", "Tải", "Chỉnh sửa"])
        for row in data.get("timeline", []):
            _append_row(ws, [
                row["period"], row["access_count"],
                row["download_count"], row["edit_count"],
            ])


# ======================= PDF (#13) =======================
def export_pdf(
    db: Session, *, user: CurrentUser, report_type: str, params: dict,
    correlation_id: Optional[str], ip: Optional[str],
) -> tuple[bytes, str]:
    _check_report_type(report_type)
    if report_type not in PDF_SUPPORTED:
        raise rc.err(
            "PDF_NOT_SUPPORTED",
            f"Báo cáo '{report_type}' chưa hỗ trợ xuất PDF (bản đầu chỉ dashboard)", 422,
        )
    _enforce_export_rbac(db, user, report_type)
    data = _gather(db, user, report_type, params)

    pdf_bytes = _render_dashboard_pdf(data, user)
    _audit_export(db, user, report_type, "pdf", params, correlation_id, ip)
    return pdf_bytes, "bao-cao-tong-quan.pdf"


# ── Template PDF dashboard ─────────────────────────────────────────────────
_ROLE_LABELS = {
    "admin": "Quản trị viên", "leader": "Ban lãnh đạo", "office": "Văn phòng",
    "staff": "KTV", "reception": "Phòng nhận mẫu", "qms": "Quản lý chất lượng",
    "lab_manager": "Trưởng phòng lab",
}
_STATUS_LABELS = {
    "received": "Đã tiếp nhận", "assigned": "Đã phân công", "testing": "Đang thực nghiệm",
    "done": "Đã chốt", "overdue": "Quá hạn", "returned": "Đã trả kết quả",
}
_MEASURE_LABELS = {"mass": "Khối lượng", "volume": "Thể tích", "count": "Đếm"}
_CHART_COLORS = ["#0d8256", "#0e7c86", "#a29f76", "#2563eb", "#ef4444", "#7c3aed"]


def _drawing_pie_samples_by_status(bs: dict, font: str):
    """Pie 'Mẫu theo trạng thái' — cùng data + màu với biểu đồ trên dashboard (Dashboard.tsx)."""
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.graphics.shapes import Drawing
    from reportlab.graphics.charts.piecharts import Pie
    from reportlab.graphics.charts.legends import Legend

    items = [(lbl, bs.get(k, 0)) for k, lbl in _STATUS_LABELS.items() if bs.get(k, 0)]
    if not items:
        return None

    d = Drawing(170 * mm, 62 * mm)
    pie = Pie()
    pie.x, pie.y = 10 * mm, 3 * mm
    pie.width = pie.height = 55 * mm
    pie.data = [v for _, v in items]
    pie.labels = None
    pie.slices.strokeWidth = 0.5
    pie.slices.strokeColor = colors.white
    for i in range(len(items)):
        pie.slices[i].fillColor = colors.HexColor(_CHART_COLORS[i % len(_CHART_COLORS)])
    d.add(pie)

    leg = Legend()
    leg.x, leg.y = 78 * mm, 50 * mm
    leg.dx = leg.dy = 6
    leg.fontName = font
    leg.fontSize = 9
    leg.alignment = "left"
    leg.columnMaximum = len(items)
    leg.colorNamePairs = [
        (colors.HexColor(_CHART_COLORS[i % len(_CHART_COLORS)]), f"{lbl} ({v})")
        for i, (lbl, v) in enumerate(items)
    ]
    d.add(leg)
    return d


def _base_vertical_bar_chart(font: str, x: float, y: float, w: float, h: float):
    """Khung VerticalBarChart dùng chung (trục/label) — mỗi caller set .data/.categoryNames/màu riêng."""
    from reportlab.graphics.charts.barcharts import VerticalBarChart

    bc = VerticalBarChart()
    bc.x, bc.y = x, y
    bc.width, bc.height = w, h
    bc.categoryAxis.labels.fontName = font
    bc.categoryAxis.labels.fontSize = 7
    bc.categoryAxis.labels.angle = 30
    bc.categoryAxis.labels.dy = -8
    bc.valueAxis.labels.fontName = font
    bc.valueAxis.labels.fontSize = 7
    return bc


def _drawing_bar_samples_over_time(points: list, font: str):
    """Bar 'Mẫu theo thời gian' — số mẫu tiếp nhận theo kỳ."""
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.graphics.shapes import Drawing

    if not points:
        return None

    cats = [p["period"] for p in points]
    vals = [p["count"] for p in points]

    d = Drawing(170 * mm, 62 * mm)
    bc = _base_vertical_bar_chart(font, 20 * mm, 15 * mm, 145 * mm, 42 * mm)
    bc.data = [vals]
    bc.categoryAxis.categoryNames = cats
    bc.valueAxis.valueMin = 0
    bc.bars[0].fillColor = colors.HexColor(_CHART_COLORS[2])
    d.add(bc)
    return d


def _drawing_bar_chemical_consumption(by_group: list, font: str):
    """Bar 'Tiêu hao hóa chất theo tháng' — 1 series/measurement_group, tách vì không cộng g + mL."""
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.graphics.shapes import Drawing
    from reportlab.graphics.charts.legends import Legend

    if not by_group:
        return None
    periods = sorted({pt["period"] for g in by_group for pt in g["data"]})
    if not periods:
        return None

    d = Drawing(170 * mm, 72 * mm)
    bc = _base_vertical_bar_chart(font, 20 * mm, 22 * mm, 145 * mm, 42 * mm)
    bc.data = [
        [next((pt["qty"] for pt in g["data"] if pt["period"] == p), 0) for p in periods]
        for g in by_group
    ]
    bc.categoryAxis.categoryNames = periods
    for i in range(len(by_group)):
        bc.bars[i].fillColor = colors.HexColor(_CHART_COLORS[i % len(_CHART_COLORS)])
    d.add(bc)

    leg = Legend()
    leg.x, leg.y = 20 * mm, 5 * mm
    leg.dx = leg.dy = 6
    leg.fontName = font
    leg.fontSize = 8
    leg.alignment = "left"
    leg.columnMaximum = 1
    leg.colorNamePairs = [
        (
            colors.HexColor(_CHART_COLORS[i % len(_CHART_COLORS)]),
            f"{_MEASURE_LABELS.get(g['measurement_group'], g['measurement_group'])} ({g['base_unit']})",
        )
        for i, g in enumerate(by_group)
    ]
    d.add(leg)
    return d


def _register_fonts():
    """Đăng ký DejaVuSans (hỗ trợ tiếng Việt); fallback Helvetica nếu thiếu."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    import os

    base = "/usr/share/fonts/truetype/dejavu"
    reg, bold = os.path.join(base, "DejaVuSans.ttf"), os.path.join(base, "DejaVuSans-Bold.ttf")
    try:
        if "VN" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("VN", reg))
            pdfmetrics.registerFont(TTFont("VN-Bold", bold))
        return "VN", "VN-Bold"
    except Exception:  # noqa: BLE001 — thiếu font → dùng mặc định (mất dấu)
        logger.warning("Không tìm thấy font DejaVuSans tại %s — PDF export sẽ mất dấu tiếng Việt", base)
        return "Helvetica", "Helvetica-Bold"


def _render_dashboard_pdf(data: dict, user: CurrentUser) -> bytes:
    from datetime import datetime, timezone, timedelta
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether,
    )

    F, FB = _register_fonts()
    INK = colors.HexColor("#1f2937")
    MUT = colors.HexColor("#6b7280")
    LINE = colors.HexColor("#d1d5db")
    HEAD_BG = colors.HexColor("#f3f4f6")
    ACCENT = colors.HexColor("#2f3a56")

    org = ParagraphStyle("org", fontName=FB, fontSize=11, textColor=ACCENT, alignment=1, leading=14)
    sub = ParagraphStyle("sub", fontName=F, fontSize=8.5, textColor=MUT, alignment=1, leading=11)
    title = ParagraphStyle("title", fontName=FB, fontSize=15, textColor=INK, alignment=1, spaceBefore=8, spaceAfter=2, leading=18)
    meta = ParagraphStyle("meta", fontName=F, fontSize=9, textColor=MUT, alignment=1, leading=12)
    h2 = ParagraphStyle("h2", fontName=FB, fontSize=11, textColor=ACCENT, spaceBefore=10, spaceAfter=4)

    scope = data.get("scope") or {}
    S = data.get("samples") or {}
    C = data.get("chemicals") or {}
    E = data.get("equipments") or {}
    H = data.get("hr") or {}
    D = data.get("documents") or {}
    N = data.get("notifications") or {}
    bs = S.get("by_status") or {}

    now = datetime.now(timezone.utc).astimezone(timezone(timedelta(hours=7)))
    story: list = []
    story.append(Paragraph("VIỆN NGHIÊN CỨU CÔNG NGHỆ SINH HỌC VÀ MÔI TRƯỜNG", org))
    story.append(Paragraph("Research Institute for Biotechnology and Environment (RIBE)", sub))
    story.append(Paragraph("BÁO CÁO TỔNG QUAN HOẠT ĐỘNG PHÒNG THÍ NGHIỆM", title))
    meta_txt = f"Người xuất: {user.full_name} — {_ROLE_LABELS.get(user.role, user.role)}"
    if scope.get("department_name"):
        meta_txt += f" · Phòng: {scope['department_name']}"
    meta_txt += f" · Xuất lúc {now.strftime('%H:%M %d/%m/%Y')}"
    story.append(Paragraph(meta_txt, meta))
    story.append(Spacer(1, 5 * mm))

    # ── Chỉ số chính ──
    kpis: list = []
    if S.get("available"):
        kpis += [("Tổng mẫu", S.get("total", 0)), ("Đang thực nghiệm", bs.get("testing", 0)),
                 ("Mẫu quá hạn", S.get("overdue", 0)), ("Đã chốt", bs.get("done", 0))]
    if C.get("available"):
        kpis += [("Hóa chất sắp hết hạn", C.get("expiring_soon", 0)),
                 ("Hóa chất tồn thấp", C.get("low_stock", 0))]
    if E.get("available"):
        kpis += [("Thiết bị quá hạn hiệu chuẩn", E.get("calibration_overdue", 0)),
                 ("Thiết bị sắp hiệu chuẩn", E.get("calibration_due_soon", 0))]
    if H.get("available"):
        kpis += [("Nâng lương / HĐ sắp tới hạn",
                  int(H.get("contract_ending", 0)) + int(H.get("salary_raise_due", 0)))]
    if D.get("available"):
        kpis += [("Tài liệu chờ duyệt", D.get("pending_review", 0))]
    if N.get("available"):
        kpis += [("Thông báo chưa đọc", N.get("unread", 0))]

    if kpis:
        story.append(Paragraph("Chỉ số chính", h2))
        # 2 cặp (nhãn | giá trị) mỗi hàng
        rows = []
        for i in range(0, len(kpis), 2):
            left = kpis[i]
            right = kpis[i + 1] if i + 1 < len(kpis) else ("", "")
            rows.append([left[0], str(left[1]), right[0], str(right[1])])
        t = Table(rows, colWidths=[52 * mm, 20 * mm, 52 * mm, 20 * mm])
        t.setStyle(TableStyle([
            ("FONT", (0, 0), (-1, -1), F, 9.5),
            ("FONT", (1, 0), (1, -1), FB, 12), ("FONT", (3, 0), (3, -1), FB, 12),
            ("TEXTCOLOR", (0, 0), (0, -1), INK), ("TEXTCOLOR", (2, 0), (2, -1), INK),
            ("TEXTCOLOR", (1, 0), (1, -1), ACCENT), ("TEXTCOLOR", (3, 0), (3, -1), ACCENT),
            ("ALIGN", (1, 0), (1, -1), "RIGHT"), ("ALIGN", (3, 0), (3, -1), "RIGHT"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LINEBELOW", (0, 0), (-1, -1), 0.4, LINE),
            ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]))
        story.append(t)

    # ── Phân bố trạng thái mẫu ──
    if S.get("available") and bs:
        srows = [["Trạng thái", "Số lượng"]]
        for k, lbl in _STATUS_LABELS.items():
            if k in bs:
                srows.append([lbl, str(bs.get(k, 0))])
        srows.append(["Tổng cộng", str(S.get("total", 0))])
        st = Table(srows, colWidths=[70 * mm, 30 * mm])
        st.setStyle(TableStyle([
            ("FONT", (0, 0), (-1, -1), F, 10),
            ("FONT", (0, 0), (-1, 0), FB, 10), ("FONT", (0, -1), (-1, -1), FB, 10),
            ("BACKGROUND", (0, 0), (-1, 0), HEAD_BG),
            ("BACKGROUND", (0, -1), (-1, -1), HEAD_BG),
            ("TEXTCOLOR", (0, 0), (-1, -1), INK),
            ("ALIGN", (1, 0), (1, -1), "CENTER"),
            ("GRID", (0, 0), (-1, -1), 0.4, LINE),
            ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))

        block: list = [Paragraph("Phân bố trạng thái mẫu", h2), st]
        pie = _drawing_pie_samples_by_status(bs, F)
        if pie is not None:
            block += [Spacer(1, 3 * mm), pie]
        story.append(KeepTogether(block))

    # ── Biểu đồ: mẫu theo thời gian / tiêu hao hóa chất ──
    charts = data.get("charts") or {}
    sot = charts.get("samples_over_time") or {}
    if sot.get("available") and sot.get("data"):
        bar1 = _drawing_bar_samples_over_time(sot["data"], F)
        if bar1 is not None:
            story.append(KeepTogether([Paragraph("Mẫu theo thời gian", h2), bar1]))

    cc = charts.get("chemical_consumption") or {}
    by_group = cc.get("by_measurement_group") or []
    if cc.get("available") and by_group:
        bar2 = _drawing_bar_chemical_consumption(by_group, F)
        if bar2 is not None:
            story.append(KeepTogether([Paragraph("Tiêu hao hóa chất theo tháng", h2), bar2]))

    def _footer(canvas, doc):
        canvas.saveState()
        canvas.setFont(F, 8)
        canvas.setFillColor(MUT)
        canvas.drawString(16 * mm, 10 * mm,
                          "Hệ thống Quản lý Phòng Thí nghiệm (LIMS) — RIBE")
        canvas.drawRightString(A4[0] - 16 * mm, 10 * mm, f"Trang {doc.page}")
        canvas.restoreState()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4, topMargin=16 * mm, bottomMargin=16 * mm,
        leftMargin=16 * mm, rightMargin=16 * mm, title="Báo cáo tổng quan",
    )
    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return buf.getvalue()
