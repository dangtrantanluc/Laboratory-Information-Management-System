"""Xuất Excel Báo cáo mẫu phân tích theo quý — đúng mẫu "Báo cáo phân tích mẫu.xlsx".

File mẫu của Văn phòng: một bảng ở cột B–G, tiêu đề cột ở dòng 1, Times New Roman 12,
viền mảnh, căn giữa — STT · MÃ SỐ MẪU · LOẠI MẪU · SỐ LƯỢNG · CHỈ TIÊU PHÂN TÍCH ·
PHÒNG NGHIÊN CỨU PHỤ TRÁCH. Cột A của file mẫu chỉ để ghi chữ "VÍ DỤ", nên để trống.

Mặc định chỉ lấy báo cáo ĐÃ CHỐT: số liệu gửi đi phải là số đã được Văn phòng xác
nhận. Sheet "Tình trạng nộp" liệt kê MỌI phòng lab kèm trạng thái, để người đọc biết
phòng nào chưa có trong bảng — thiếu dòng mà không nói ra là cách nhanh nhất làm sai
một báo cáo nộp lên trên.
"""
import io
from datetime import datetime
from typing import Optional

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.concurrency import export_slot
from app.core.deps import CurrentUser
from app.core.exceptions import forbidden
from app.models.department import Department
from app.models.quarterly_sample_report import (
    REPORT_STATUS_LABELS, QuarterlyReport, QuarterlyReportItem,
)
from app.services import audit_service
from app.services import quarterly_sample_report_service as svc

_FONT = "Times New Roman"
# (tiêu đề, độ rộng) — cột B..G; độ rộng lấy từ file mẫu, B để mặc định như file mẫu.
_COLUMNS: list[tuple[str, Optional[float]]] = [
    ("STT", None),
    ("MÃ SỐ MẪU", 15.0),
    ("LOẠI MẪU", 18.63),
    ("SỐ LƯỢNG", 13.09),
    # File mẫu rộng 20 vì ví dụ chỉ có "Tro, độ ẩm"; danh sách chỉ tiêu thật dài hơn nhiều,
    # để 20 thì mỗi dòng cao 4–5 dòng chữ.
    ("CHỈ TIÊU PHÂN TÍCH", 40.0),
    ("PHÒNG NGHIÊN CỨU PHỤ TRÁCH", 42.54),
]
_FIRST_COL = 2  # cột B
_ROW_HEIGHT = 15.5


def _styles():
    from openpyxl.styles import Alignment, Border, Font, Side

    thin = Side(style="thin")
    return {
        "border": Border(left=thin, right=thin, top=thin, bottom=thin),
        "center": Alignment(horizontal="center", vertical="center", wrap_text=True),
        "header": Alignment(horizontal="center", vertical="center"),
        "left": Alignment(horizontal="left", vertical="center", wrap_text=True),
        "body": Font(name=_FONT, size=12),
        "bold": Font(name=_FONT, size=12, bold=True),
        "title": Font(name=_FONT, size=14, bold=True),
    }


def _print_setup(ws, title_rows: str) -> None:
    """In ngang, vừa một trang bề ngang, lặp lại hàng tiêu đề cột ở mỗi trang."""
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = title_rows


def _put(ws, row: int, col: int, value, font, align, border):
    c = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str):
        c.data_type = "s"  # văn bản thuần — chuỗi bắt đầu bằng '=' không thành công thức
    c.font, c.alignment, c.border = font, align, border
    return c


def build_workbook(
    *, year: int, quarter: int, rows: list[dict], status_rows: list[dict],
) -> bytes:
    """rows: sample_code, matrix_label, quantity, criteria_text, department_name.
    status_rows: department_name, status_label, item_count, total_quantity,
    submitted_at, locked_at, included."""
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    st = _styles()
    wb = Workbook()
    ws = wb.active
    ws.title = f"QUÝ {svc.ROMAN[quarter]}.{year}"
    for i, (label, width) in enumerate(_COLUMNS):
        col = _FIRST_COL + i
        if width:
            ws.column_dimensions[get_column_letter(col)].width = width
        _put(ws, 1, col, label, st["bold"], st["header"], st["border"])
    ws.row_dimensions[1].height = _ROW_HEIGHT
    _print_setup(ws, "1:1")

    for n, r in enumerate(rows, start=1):
        values = [
            n, r["sample_code"], r["matrix_label"], r["quantity"], r["criteria_text"],
            r["department_name"],
        ]
        for i, v in enumerate(values):
            _put(ws, n + 1, _FIRST_COL + i, v, st["body"], st["center"], st["border"])

    # ── Sheet tình trạng nộp ──
    ss = wb.create_sheet("Tình trạng nộp")
    ss.cell(row=1, column=1, value=(
        f"TÌNH TRẠNG NỘP BÁO CÁO MẪU PHÂN TÍCH QUÝ {svc.ROMAN[quarter]}/{year}"
    )).font = st["title"]
    ss.merge_cells(start_row=1, start_column=1, end_row=1, end_column=8)
    ss.cell(row=1, column=1).alignment = st["center"]
    headers = [
        ("STT", 6), ("PHÒNG NGHIÊN CỨU", 45), ("TRẠNG THÁI", 16), ("SỐ DÒNG MẪU", 12),
        ("TỔNG SỐ LƯỢNG", 12), ("NỘP LÚC", 18), ("CHỐT LÚC", 18), ("CÓ TRONG BẢNG TỔNG HỢP", 22),
    ]
    for i, (label, width) in enumerate(headers, start=1):
        ss.column_dimensions[get_column_letter(i)].width = width
        _put(ss, 3, i, label, st["bold"], st["center"], st["border"])
    ss.row_dimensions[3].height = 31
    _print_setup(ss, "3:3")

    def fmt(dt: Optional[datetime]) -> str:
        return dt.astimezone(svc.VN_TZ).strftime("%d/%m/%Y %H:%M") if dt else ""

    for n, s in enumerate(status_rows, start=1):
        values = [
            n, s["department_name"], s["status_label"], s["item_count"], s["total_quantity"],
            fmt(s["submitted_at"]), fmt(s["locked_at"]), "Có" if s["included"] else "Không",
        ]
        for i, v in enumerate(values, start=1):
            _put(ss, n + 3, i, v, st["body"], st["left"] if i == 2 else st["center"], st["border"])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_xlsx(
    db: Session, *, user: CurrentUser, year: int, quarter: int, include_unlocked: bool,
    correlation_id: Optional[str], ip: Optional[str],
) -> tuple[bytes, str]:
    if not svc.can_read_all(db, user):
        raise forbidden("Chỉ Văn phòng / Ban lãnh đạo xuất được báo cáo toàn Viện")
    with export_slot():
        depts = db.execute(
            select(Department)
            .where(Department.kind == "lab", Department.status == "active")
            .order_by(Department.name)
        ).scalars().all()
        reports = {
            r.department_id: r
            for r in db.execute(
                select(QuarterlyReport).where(
                    QuarterlyReport.year == year, QuarterlyReport.quarter == quarter,
                )
            ).scalars()
        }

        rows: list[dict] = []
        status_rows: list[dict] = []
        for d in depts:
            r = reports.get(d.id)
            items = svc.load_items(db, r.id) if r else []
            included = r is not None and (include_unlocked or r.status == "locked")
            status_rows.append({
                "department_name": d.name,
                "status_label": REPORT_STATUS_LABELS[r.status] if r else "Chưa bắt đầu",
                "item_count": len(items),
                "total_quantity": sum(i.quantity for i in items),
                "submitted_at": r.submitted_at if r else None,
                "locked_at": r.locked_at if r else None,
                "included": included,
            })
            if not included:
                continue
            crits = svc.criteria_by_item(db, [i.id for i in items])
            rows.extend(_row(i, crits[i.id], d.name) for i in items)

        content = build_workbook(year=year, quarter=quarter, rows=rows, status_rows=status_rows)
        audit_service.log_action(
            db, action="QUARTERLY_SAMPLE_REPORT_EXPORT", resource="quarterly_sample_report",
            user_id=user.id, resource_id=None, correlation_id=correlation_id, ip=ip,
            detail={
                "year": year, "quarter": quarter, "rows": len(rows),
                "include_unlocked": include_unlocked,
            },
        )
        db.commit()
    return content, f"bao-cao-mau-quy-{svc.ROMAN[quarter]}-{year}.xlsx"


def _row(it: QuarterlyReportItem, criteria: list, dept_name: str) -> dict:
    return {
        "sample_code": it.sample_code,
        "matrix_label": svc.matrix_label(it.matrix, it.matrix_other),
        "quantity": it.quantity,
        "criteria_text": svc.criteria_text(criteria, it.extra_criteria),
        "department_name": dept_name,
    }
